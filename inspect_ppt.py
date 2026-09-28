import sys
from pptx import Presentation

prs = Presentation(sys.argv[1])
print(f"Original slide count: {len(prs.slides)}")

# Let's inspect slide titles
for i, slide in enumerate(prs.slides):
    title = slide.shapes.title.text if slide.shapes.title else "No Title"
    print(f"Slide {i+1}: {title}")
