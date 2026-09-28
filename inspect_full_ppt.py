import sys
from pptx import Presentation

prs = Presentation(sys.argv[1])

# Let's inspect all text frames and shapes in each slide to understand how we can append or update.
for idx, slide in enumerate(prs.slides):
    print(f"=== SLIDE {idx+1} ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            print(f"[Shape ID: {shape.shape_id}] Text:")
            print(shape.text_frame.text)
            print("-" * 40)
