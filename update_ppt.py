import sys
from pptx import Presentation

def add_slide(prs, title, content_bullets):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.text = content_bullets[0]
    
    for bullet in content_bullets[1:]:
        p = tf.add_paragraph()
        p.text = bullet
        p.level = 0

prs = Presentation(sys.argv[1])
# We can add new slides or inspect
print(f"Total slides: {len(prs.slides)}")
