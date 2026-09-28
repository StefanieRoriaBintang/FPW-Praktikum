import sys
from pptx import Presentation

prs = Presentation(sys.argv[1])
for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} shapes ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            print(shape.text_frame.text[:100])
