import sys
from pptx import Presentation
from pptx.util import Inches, Pt

input_path = sys.argv[1]
output_path = sys.argv[2]

prs = Presentation(input_path)
layout = prs.slide_layouts[0] # DEFAULT layout

def add_detailed_slide(prs, title_text, items):
    slide = prs.slides.add_slide(layout)
    
    # Since layout has custom shapes or we can add textbox
    # Let's add title and text box manually for reliability
    txBox = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.0))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(24)
    
    txBox_content = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.5))
    tf_content = txBox_content.text_frame
    tf_content.word_wrap = True
    
    for idx, (subtitle, desc) in enumerate(items):
        p_sub = tf_content.add_paragraph() if idx > 0 else tf_content.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.bold = True
        p_sub.font.size = Pt(14)
        
        p_desc = tf_content.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(12)
        p_desc.font.name = 'Consolas'

# Add Slide for Complete Eloquent CRUD & Controller
add_detailed_slide(prs, "IMPLEMENTASI: Eloquent CRUD di Controller", [
    ("1. Create (Simpan Data)", "Post::create(['title' => $req->title, 'body' => $req->body, 'user_id' => $req->user_id]);"),
    ("2. Read (Ambil Data)", "$posts = Post::with('user')->get(); // Mengambil data beserta relasi user"),
    ("3. Update (Ubah Data)", "$post->update(['title' => $req->title]); // Memperbarui data post"),
    ("4. Delete (Hapus Data)", "$post->delete(); // Menghapus data dari database")
])

# Add Slide for Model Relationships Implementation
add_detailed_slide(prs, "IMPLEMENTASI: Model Relationships", [
    ("One to Many (User -> Posts)", "Di User.php: return $this->hasMany(Post::class);\nDi Post.php: return $this->belongsTo(User::class);"),
    ("One to One (User -> Profile)", "Di User.php: return $this->hasOne(Profile::class);\nDi Profile.php: return $this->belongsTo(User::class);"),
    ("Many to Many (Post <-> Tag)", "Di Post.php: return $this->belongsToMany(Tag::class);\nDi Tag.php: return $this->belongsToMany(Post::class);")
])

# Add Slide for Query Builder Practical Code
add_detailed_slide(prs, "IMPLEMENTASI: Query Builder (DB Facade)", [
    ("SELECT & WHERE", "$posts = DB::table('posts')->where('user_id', 1)->get();"),
    ("INSERT", "DB::table('posts')->insert(['title' => 'Judul', 'body' => 'Isi', 'user_id' => 1]);"),
    ("UPDATE & DELETE", "DB::table('posts')->where('id', 1)->update(['title' => 'Baru']);\nDB::table('posts')->where('id', 1)->delete();")
])

prs.save(output_path)
print(f"Successfully updated presentation saved to {output_path}")
