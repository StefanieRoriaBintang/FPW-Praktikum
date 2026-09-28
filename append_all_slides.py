import sys
from pptx import Presentation
from pptx.util import Inches, Pt

input_path = sys.argv[1]
output_path = sys.argv[2]

prs = Presentation(input_path)
layout = prs.slide_layouts[0] # DEFAULT layout

def add_detailed_slide(prs, title_text, items):
    slide = prs.slides.add_slide(layout)
    
    txBox = slide.shapes.add_textbox(Inches(1.0), Inches(0.5), Inches(11.3), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(22)
    
    txBox_content = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(5.2))
    tf_content = txBox_content.text_frame
    tf_content.word_wrap = True
    
    for idx, (subtitle, desc) in enumerate(items):
        p_sub = tf_content.add_paragraph() if idx > 0 else tf_content.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.bold = True
        p_sub.font.size = Pt(13)
        
        p_desc = tf_content.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.name = 'Consolas'

# 1. Database Migration
add_detailed_slide(prs, "IMPLEMENTASI DETAIL: Database Migration", [
    ("1. Membuat File Migration", "php artisan make:migration create_posts_table\nDigunakan untuk membuat file skema tabel baru di database/migrations."),
    ("2. Menulis Struktur Tabel (up method)", "Schema::create('posts', function (Blueprint $table) {\n    $table->id();\n    $table->string('title');\n    $table->text('body');\n    $table->foreignId('user_id')->constrained()->cascadeOnDelete();\n    $table->timestamps();\n});"),
    ("3. Perintah Artisan Migration", "• php artisan migrate (menjalankan migration)\n• php artisan migrate:fresh --seed (reset semua tabel dan jalankan seeder)")
])

# 2. Seeder & Factory
add_detailed_slide(prs, "IMPLEMENTASI DETAIL: Seeder & Factory", [
    ("1. Membuat Factory", "php artisan make:factory PostFactory --model=Post\nMenghasilkan data dummy realistis menggunakan Faker."),
    ("2. Konfigurasi Factory (PostFactory.php)", "public function definition(): array {\n    return [\n        'title' => fake()->sentence(),\n        'body' => fake()->paragraphs(3, true),\n        'user_id' => User::factory(),\n    ];\n}"),
    ("3. Membuat & Menjalankan Seeder", "php artisan make:seeder PostSeeder\nDi run(): User::factory(5)->has(Post::factory(10))->create();\nJalankan dengan: php artisan db:seed")
])

# 3. Eloquent ORM
add_detailed_slide(prs, "IMPLEMENTASI DETAIL: Eloquent ORM", [
    ("1. Definisi Model & Mass Assignment", "class Post extends Model {\n    use HasFactory;\n    protected $fillable = ['title', 'body', 'user_id']; // Kolom yang diizinkan diisi massal\n}"),
    ("2. Operasi CRUD Eloquent", "• CREATE: Post::create(['title' => 'Judul', 'body' => 'Isi', 'user_id' => 1]);\n• READ: Post::all(); atau Post::find($id);\n• UPDATE: $post->update(['title' => 'Judul Baru']);\n• DELETE: $post->delete();")
])

# 4. Model Relationships
add_detailed_slide(prs, "IMPLEMENTASI DETAIL: Model Relationships", [
    ("1. One to One", "User hasOne Profile -> public function profile() { return $this->hasOne(Profile::class); }\nProfile belongsTo User -> public function user() { return $this->belongsTo(User::class); }"),
    ("2. One to Many", "User hasMany Post -> public function posts() { return $this->hasMany(Post::class); }\nPost belongsTo User -> public function user() { return $this->belongsTo(User::class); }"),
    ("3. Many to Many", "Post belongsToMany Tag -> return $this->belongsToMany(Tag::class); (via pivot table post_tag)")
])

# 5. Query Builder
add_detailed_slide(prs, "IMPLEMENTASI DETAIL: Query Builder (DB Facade)", [
    ("1. Mengambil Data (Select & Where)", "$posts = DB::table('posts')->where('user_id', 1)->get();"),
    ("2. Menyimpan Data (Insert)", "DB::table('posts')->insert(['title' => 'Judul', 'body' => 'Isi', 'user_id' => 1, 'created_at' => now(), 'updated_at' => now()]);"),
    ("3. Mengubah & Menghapus (Update & Delete)", "DB::table('posts')->where('id', 1)->update(['title' => 'Baru']);\nDB::table('posts')->where('id', 1)->delete();")
])

prs.save(output_path)
print(f"Successfully updated presentation saved to {output_path}")
