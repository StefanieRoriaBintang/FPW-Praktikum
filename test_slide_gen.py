import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation(sys.argv[1])

def add_content_slide(prs, title_text, sections):
    """
    sections is a list of tuples: (heading, code_snippet_or_text)
    """
    blank_layout = prs.slide_layouts[6] # Blank slide or title/content
    slide = prs.slides.add_slide(prs.slide_layouts[1]) # Use Title and Content layout
    
    # Set title
    title_shape = slide.shapes.title
    title_shape.text = title_text
    
    # Get body shape / text frame
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.clear()
    
    for i, (heading, details) in enumerate(sections):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = heading
        p.font.bold = True
        p.font.size = Pt(14)
        
        p2 = tf.add_paragraph()
        p2.text = details
        p2.font.size = Pt(12)
        p2.font.name = 'Consolas'

print("Helper ready")
